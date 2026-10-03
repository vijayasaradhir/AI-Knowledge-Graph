from __future__ import annotations

import json
import os
import re
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Sequence

from .code_graph import CodeGraphIngestor
from .embeddings import TfidfEmbeddingEngine
from .extractor import EntityRelationExtractor
from .graph_store import KnowledgeGraph, Neo4jGraphStore
from .models import Document, Entity, ExtractionResult, Relation, QueryResult
from .reasoning import GraphReasoner

try:  # pragma: no cover - optional dependency
    from openai import OpenAI
except Exception:  # pragma: no cover
    OpenAI = None


_TEXT_FIELDS = {"text", "content", "body", "description", "summary"}
_SUPPORTED_TEXT_EXTENSIONS = {".txt", ".md", ".json", ".log"}
_SUPPORTED_BINARY_EXTENSIONS = {".pdf", ".docx", ".doc"}


@dataclass(slots=True)
class _RagContext:
    query: str
    answer_hint: str
    matched_entities: List[Entity]
    supporting_relations: List[Relation]
    documents: List[Document]


class SmartAssistPipeline:
    def __init__(self, enable_code_graph: Optional[bool] = None) -> None:
        self.extractor = EntityRelationExtractor()
        self.embeddings = TfidfEmbeddingEngine()
        self.graph = KnowledgeGraph()
        self.documents: List[Document] = []
        self.code_graph = CodeGraphIngestor(enabled=self._resolve_code_graph_enabled(enable_code_graph))
        self._openai_client = self._build_openai_client()
        self._answer_model = os.getenv("KG_OPENAI_ANSWER_MODEL", os.getenv("KG_OPENAI_MODEL", "gpt-4.1-mini"))
        self.neo4j_store = self._build_neo4j_store()
        self.neo4j_enabled = self.neo4j_store is not None

    def _resolve_code_graph_enabled(self, explicit: Optional[bool]) -> bool:
        if explicit is not None:
            return explicit
        value = os.getenv("SMART_ASSIST_CODE_GRAPH", os.getenv("KG_CODE_GRAPH", "0"))
        return value.strip().lower() in {"1", "true", "yes", "on"}

    def _build_neo4j_store(self) -> Optional[Neo4jGraphStore]:
        uri = os.getenv("NEO4J_URI")
        user = os.getenv("NEO4J_USER")
        password = os.getenv("NEO4J_PASSWORD")
        database = os.getenv("NEO4J_DATABASE", "neo4j")
        if uri and user and password:
            try:
                return Neo4jGraphStore(uri, user, password, database=database)
            except Exception as exc:
                warnings.warn(f"Neo4j is configured but unavailable: {exc}", RuntimeWarning)
                return None
        return None

    def load_documents(self, source: object, source_type: str = "auto") -> List[Document]:
        normalized_source_type = source_type.strip().lower()
        if normalized_source_type == "code":
            if not self.code_graph.enabled:
                raise RuntimeError(
                    "Code graph ingestion is disabled. Enable SMART_ASSIST_CODE_GRAPH or pass enable_code_graph=True."
                )
            return self.code_graph.load_documents(source)

        if isinstance(source, (str, Path)):
            path = Path(source)
            if path.is_dir():
                documents = self._load_documents_from_directory(path)
                if self.code_graph.enabled:
                    documents.extend(self.code_graph.load_documents(path))
                return documents
            if self.code_graph.enabled and path.suffix.lower() in {".py", ".java"}:
                return self.code_graph.load_documents(path)
            return self._load_documents_from_file(path)

        if isinstance(source, dict):
            data = [source]
        else:
            data = list(source)
        documents: List[Document] = []
        for index, row in enumerate(data):
            if isinstance(row, dict):
                documents.append(
                    Document(
                        id=str(row.get("id") or f"doc-{index+1:03d}"),
                        text=self._document_text_from_record(row),
                        metadata={k: v for k, v in row.items() if k not in {"id", *_TEXT_FIELDS}},
                    )
                )
                continue
            documents.append(
                Document(
                    id=f"doc-{index+1:03d}",
                    text=str(row),
                    metadata={},
                )
            )
        return documents

    def _load_documents_from_directory(self, directory: Path) -> List[Document]:
        files = sorted(
            path
            for path in directory.rglob("*")
            if path.is_file() and path.suffix.lower() in _SUPPORTED_TEXT_EXTENSIONS | _SUPPORTED_BINARY_EXTENSIONS
        )
        if not files:
            raise ValueError(f"No supported documents found in directory: {directory}")
        documents: List[Document] = []
        for file_path in files:
            documents.extend(self._load_documents_from_file(file_path, root=directory))
        return documents

    def _load_documents_from_file(self, path: Path, root: Optional[Path] = None) -> List[Document]:
        suffix = path.suffix.lower()
        if suffix not in _SUPPORTED_TEXT_EXTENSIONS | _SUPPORTED_BINARY_EXTENSIONS:
            raise ValueError(f"Unsupported document format: {path.suffix}")

        if suffix in _SUPPORTED_TEXT_EXTENSIONS:
            return [self._document_from_text(path, root=root)]

        if suffix == ".pdf":
            return [self._document_from_extracted_text(path, self._extract_text_from_pdf(path), root=root)]

        if suffix == ".docx":
            return [self._document_from_extracted_text(path, self._extract_text_from_docx(path), root=root)]

        if suffix == ".doc":
            return [self._document_from_extracted_text(path, self._extract_text_from_doc(path), root=root)]

        raw = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(raw, dict):
            raw = [raw]
        if not isinstance(raw, list):
            raise ValueError(f"JSON file must contain an object or an array of objects: {path}")
        if not raw:
            raise ValueError(f"JSON file contains no documents: {path}")

        documents: List[Document] = []
        for index, row in enumerate(raw):
            if not isinstance(row, dict):
                raise ValueError(f"Each JSON record must be an object: {path}")
            documents.append(
                Document(
                    id=str(row.get("id") or self._document_id_from_path(path, root=root, index=index)),
                    text=self._document_text_from_record(row),
                    metadata={
                        **{k: v for k, v in row.items() if k not in {"id", *_TEXT_FIELDS}},
                        "source_file": str(path),
                        "format": "json",
                    },
                )
            )
        return documents

    def _document_text_from_record(self, row: dict) -> str:
        for field in ("text", "content", "body", "description", "summary"):
            value = row.get(field)
            if value is not None:
                text = str(value).strip()
                if text:
                    return text

        if len(row) == 1:
            value = next(iter(row.values()))
            if isinstance(value, str) and value.strip():
                return value.strip()

        return json.dumps(row, ensure_ascii=False, indent=2, sort_keys=True)

    def _document_from_text(self, path: Path, root: Optional[Path] = None) -> Document:
        return Document(
            id=self._document_id_from_path(path, root=root),
            text=path.read_text(encoding="utf-8"),
            metadata={
                "source_file": str(path),
                "format": path.suffix.lower().lstrip("."),
            },
        )

    def _document_from_extracted_text(self, path: Path, text: str, root: Optional[Path] = None) -> Document:
        cleaned = self._normalize_extracted_text(text)
        if not cleaned:
            raise ValueError(f"No extractable text found in document: {path}")
        return Document(
            id=self._document_id_from_path(path, root=root),
            text=cleaned,
            metadata={
                "source_file": str(path),
                "format": path.suffix.lower().lstrip("."),
                "extracted": True,
            },
        )

    def _extract_text_from_pdf(self, path: Path) -> str:
        try:
            from pypdf import PdfReader
        except Exception as exc:
            raise RuntimeError(
                "PDF support requires the optional 'pypdf' dependency. Install it to ingest .pdf files."
            ) from exc

        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pages.append("")
        return "\n".join(part for part in pages if part.strip())

    def _extract_text_from_docx(self, path: Path) -> str:
        try:
            from docx import Document as DocxDocument
        except Exception as exc:
            raise RuntimeError(
                "DOCX support requires the optional 'python-docx' dependency. Install it to ingest .docx files."
            ) from exc

        doc = DocxDocument(str(path))
        parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)

    def _extract_text_from_doc(self, path: Path) -> str:
        try:
            import win32com.client  # type: ignore
        except Exception as exc:
            raise RuntimeError(
                "Legacy .doc support requires Microsoft Word and the optional 'pywin32' dependency on Windows. "
                "If you cannot install that, convert the file to .docx before ingesting."
            ) from exc

        word = None
        document = None
        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            document = word.Documents.Open(str(path), ReadOnly=True)
            text = document.Content.Text
            return str(text)
        except Exception as exc:
            raise RuntimeError(f"Unable to extract text from legacy Word document: {path}") from exc
        finally:
            if document is not None:
                try:
                    document.Close(False)
                except Exception:
                    pass
            if word is not None:
                try:
                    word.Quit()
                except Exception:
                    pass

    def _document_id_from_path(self, path: Path, root: Optional[Path] = None, index: Optional[int] = None) -> str:
        if root is not None:
            try:
                relative = path.relative_to(root)
            except ValueError:
                relative = path.name
            base = str(relative).replace("\\", "/")
        else:
            base = path.name
        normalized = base.replace("/", "__")
        if index is not None:
            return f"{normalized}::{index+1:03d}"
        return normalized

    def ingest(self, documents: Sequence[Document]) -> List[ExtractionResult]:
        if not documents:
            raise ValueError("No documents were loaded for ingestion.")
        self.documents = list(documents)
        texts = [doc.text for doc in documents]
        self.embeddings.fit(texts)

        results: List[Optional[ExtractionResult]] = [None] * len(documents)
        code_documents: list[tuple[int, Document]] = []

        for index, document in enumerate(documents):
            if self.code_graph.should_handle(document):
                code_documents.append((index, document))
                continue
            self.graph.add_document(document)
            extracted = self.extractor.extract(document)
            results[index] = extracted

            for entity in extracted.entities:
                self.graph.add_entity(entity)
                self.graph.add_mention(document.id, entity.id)

            for relation in extracted.relations:
                self.graph.add_relation(relation)

        if code_documents:
            code_results = self.code_graph.ingest_documents([document for _, document in code_documents], self.graph)
            for (index, _), result in zip(code_documents, code_results):
                results[index] = result

        self._add_similarity_edges(documents)
        if self.neo4j_store:
            try:
                self.neo4j_store.sync(self.graph)
            except Exception as exc:
                warnings.warn(f"Neo4j sync failed; graph was not persisted: {exc}", RuntimeWarning)
        return [result for result in results if result is not None]

    def _add_similarity_edges(self, documents: Sequence[Document]) -> None:
        for i, left in enumerate(documents):
            for j in range(i + 1, len(documents)):
                right = documents[j]
                left_kind = str(left.metadata.get("kind", "text")).lower()
                right_kind = str(right.metadata.get("kind", "text")).lower()
                if left_kind != right_kind:
                    continue
                if left_kind == "code" and str(left.metadata.get("language", "")).lower() != str(
                    right.metadata.get("language", "")
                ).lower():
                    continue
                score = self.embeddings.similarity(left.text, right.text)
                if score >= 0.18:
                    self.graph.add_similarity(left.id, right.id, score, "document_embeddings")
                    self.graph.add_similarity(right.id, left.id, score, "document_embeddings")

    def reason(self, text: str, top_k: int = 5) -> QueryResult:
        self._ensure_loaded_for_query()
        graph_reasoner = GraphReasoner(self.graph)
        exact = graph_reasoner.query(text, top_k=top_k)
        if not self.documents:
            return exact

        doc_texts = [doc.text for doc in self.documents]
        matches = self.embeddings.nearest_neighbors(text, doc_texts, top_k=top_k)
        rag_context = self._build_rag_context(text, exact, matches, top_k=top_k)
        llm_result = self._answer_with_llm(rag_context)
        if llm_result is None:
            return self._enrich_graph_answer(exact, matches, top_k=top_k)

        return QueryResult(
            matched_entities=llm_result.get("matched_entities", exact.matched_entities)[:top_k],
            similar_entities=llm_result.get("similar_entities", exact.similar_entities)[:top_k],
            supporting_relations=llm_result.get("supporting_relations", exact.supporting_relations)[:top_k],
            explanation=llm_result.get("explanation", exact.explanation),
            answer=llm_result.get("answer", exact.answer),
            intent=exact.intent,
        )

    def _build_openai_client(self) -> Optional["OpenAI"]:
        if OpenAI is None:
            return None
        api_key = os.getenv("OPENAI_API_KEY")
        if not api_key:
            return None
        try:
            return OpenAI(api_key=api_key)
        except Exception:
            return None

    def _ensure_loaded_for_query(self) -> None:
        if self.documents or not self.neo4j_store:
            return
        try:
            documents = self.neo4j_store.load(self.graph)
        except Exception as exc:
            warnings.warn(f"Neo4j load failed; querying in-memory graph only: {exc}", RuntimeWarning)
            return
        if not documents:
            return
        self.documents = documents
        self.embeddings.fit([doc.text for doc in documents])

    def _build_rag_context(
        self,
        query: str,
        exact: QueryResult,
        matches,
        top_k: int = 5,
    ) -> _RagContext:
        related_entities: list[str] = []
        related_relations: list[Relation] = list(exact.supporting_relations)
        seen_related_entities: set[str] = set()

        for match in matches:
            document = self.documents[match.left_index]
            for node_id, _ in self._entities_for_document(document.id):
                if node_id not in seen_related_entities:
                    seen_related_entities.add(node_id)
                    related_entities.append(node_id)
                for source, target, _, relation_data in self.graph.graph.out_edges(node_id, keys=True, data=True):
                    if relation_data.get("kind") == "relation":
                        related_relations.append(self._relation_from_edge(source, target, relation_data))

        matched_entities: dict[str, Entity] = {entity.id: entity for entity in exact.matched_entities}
        for entity_id in related_entities:
            node = self.graph.graph.nodes[entity_id] if entity_id in self.graph.graph.nodes else {}
            if node.get("kind") != "entity":
                continue
            if entity_id not in matched_entities:
                matched_entities[entity_id] = self._entity_from_node(entity_id, node)

        docs = [self.documents[match.left_index] for match in matches[:top_k]]
        return _RagContext(
            query=query,
            answer_hint=exact.answer,
            matched_entities=list(matched_entities.values())[:top_k],
            supporting_relations=self._dedupe_relations(related_relations)[:top_k],
            documents=docs,
        )

    def _answer_with_llm(self, context: _RagContext) -> Optional[dict]:
        if self._openai_client is None:
            return None

        schema = {
            "name": "smart_assist_query_answer",
            "schema": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "answer": {"type": "string"},
                    "explanation": {"type": "string"},
                    "used_documents": {"type": "array", "items": {"type": "string"}},
                    "used_entities": {"type": "array", "items": {"type": "string"}},
                    "used_relations": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["answer", "explanation", "used_documents", "used_entities", "used_relations"],
            },
            "strict": True,
        }

        system_prompt = (
            "You answer user questions using only the provided graph and document context.\n"
            "If the context is insufficient, say so clearly instead of inventing facts.\n"
            "Return concise, grounded answers."
        )
        user_prompt = self._build_query_prompt(context)

        try:
            response = self._openai_client.chat.completions.create(
                model=self._answer_model,
                temperature=0,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                response_format={"type": "json_schema", "json_schema": schema},
            )
            content = response.choices[0].message.content or "{}"
            payload = json.loads(content)
            payload["matched_entities"] = context.matched_entities
            payload["supporting_relations"] = context.supporting_relations
            payload["similar_entities"] = []
            return payload
        except Exception as exc:
            warnings.warn(f"LLM query failed; falling back to graph reasoning: {exc}", RuntimeWarning)
            return None

    def _build_query_prompt(self, context: _RagContext) -> str:
        lines = [
            "Answer the user query using the grounded context below.",
            "",
            f"Query: {context.query}",
            "",
            "Graph answer hint:",
            context.answer_hint or "No direct graph answer was available.",
            "",
            "Matched entities:",
        ]
        if context.matched_entities:
            for entity in context.matched_entities:
                aliases = ", ".join(entity.aliases) if entity.aliases else "none"
                lines.append(
                    f"- {entity.name} | label={entity.label} | confidence={entity.confidence:.2f} | aliases={aliases}"
                )
        else:
            lines.append("- none")

        lines.append("")
        lines.append("Supporting relations:")
        if context.supporting_relations:
            for relation in context.supporting_relations:
                lines.append(
                    f"- {relation.source_id} -> {relation.target_id} | type={relation.relation_type} | evidence={relation.evidence}"
                )
        else:
            lines.append("- none")

        lines.append("")
        lines.append("Retrieved documents:")
        if context.documents:
            for document in context.documents:
                lines.append(f"- {document.id}: {self._document_excerpt(document.text, context.query)}")
        else:
            lines.append("- none")

        lines.append("")
        lines.append("Rules:")
        lines.append("- Use only the context above.")
        lines.append("- Mention when evidence is weak or missing.")
        lines.append("- Keep the answer concise and practical.")
        return "\n".join(lines)

    def _document_excerpt(self, text: str, query: str, max_chars: int = 500) -> str:
        cleaned = " ".join(text.split())
        if not cleaned:
            return ""

        lowered = cleaned.lower()
        query_terms = [term for term in re.findall(r"[A-Za-z0-9_]{3,}", query.lower()) if term not in {"the", "and", "for", "with", "from", "that", "this"}]
        for term in query_terms:
            idx = lowered.find(term)
            if idx != -1:
                start = max(0, idx - 180)
                end = min(len(cleaned), idx + 320)
                excerpt = cleaned[start:end].strip()
                if start > 0:
                    excerpt = "..." + excerpt
                if end < len(cleaned):
                    excerpt = excerpt + "..."
                return excerpt[:max_chars]
        return (cleaned[: max_chars - 3] + "...") if len(cleaned) > max_chars else cleaned

    def _normalize_extracted_text(self, text: str) -> str:
        cleaned = text.replace("\x00", " ")
        cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
        cleaned = re.sub(r"[ \t]+", " ", cleaned)
        return cleaned.strip()

    def _enrich_graph_answer(self, exact: QueryResult, matches, top_k: int = 5) -> QueryResult:
        related_entities: list[str] = []
        related_relations: list[Relation] = list(exact.supporting_relations)
        seen_related_entities: set[str] = set()

        for match in matches:
            document = self.documents[match.left_index]
            for node_id, _ in self._entities_for_document(document.id):
                if node_id not in seen_related_entities:
                    seen_related_entities.add(node_id)
                    related_entities.append(node_id)
                for source, target, _, relation_data in self.graph.graph.out_edges(node_id, keys=True, data=True):
                    if relation_data.get("kind") == "relation":
                        related_relations.append(self._relation_from_edge(source, target, relation_data))

        matched_entities: dict[str, Entity] = {entity.id: entity for entity in exact.matched_entities}
        for entity_id in related_entities:
            node = self.graph.graph.nodes[entity_id] if entity_id in self.graph.graph.nodes else {}
            if node.get("kind") != "entity":
                continue
            if entity_id not in matched_entities:
                matched_entities[entity_id] = self._entity_from_node(entity_id, node)

        similar_entities = list(exact.similar_entities)
        seen_similar = {entity.id for entity in similar_entities}
        for entity_id in related_entities:
            if entity_id in seen_similar:
                continue
            node = self.graph.graph.nodes[entity_id] if entity_id in self.graph.graph.nodes else {}
            if node.get("kind") == "entity":
                similar_entities.append(self._entity_from_node(entity_id, node))
                seen_similar.add(entity_id)

        explanation_bits = [exact.explanation]
        if matches:
            explanation_bits.append(
                "Embedding search surfaced " + ", ".join(f"{self.documents[m.left_index].id} ({m.score:.2f})" for m in matches)
            )
        fallback_answer = exact.answer
        if fallback_answer.startswith("No direct answer") and related_entities:
            fallback_answer = GraphReasoner(self.graph).compose_answer(
                list(matched_entities.values())[:top_k],
                self._dedupe_relations(related_relations)[:top_k],
                exact.intent,
            )
        return QueryResult(
            matched_entities=list(matched_entities.values())[:top_k],
            similar_entities=similar_entities[:top_k],
            supporting_relations=self._dedupe_relations(related_relations)[:top_k],
            explanation=" ".join(explanation_bits),
            answer=fallback_answer,
            intent=exact.intent,
        )

    def _entities_for_document(self, document_id: str):
        for _, target, _, data in self.graph.graph.out_edges(document_id, keys=True, data=True):
            if data.get("kind") == "mentions":
                yield target, data

    def _dedupe_relations(self, relations: List[Relation]) -> List[Relation]:
        seen = set()
        deduped: List[Relation] = []
        for relation in relations:
            key = (relation.source_id, relation.target_id, relation.relation_type, relation.evidence)
            if key in seen:
                continue
            seen.add(key)
            deduped.append(relation)
        return deduped

    def _entity_from_node(self, entity_id: str, node: dict):
        return Entity(
            id=entity_id,
            name=node.get("name", ""),
            label=node.get("label", "Concept"),
            source_document_id=node.get("metadata", {}).get("documents", [""])[0] if node.get("metadata") else "",
            confidence=node.get("confidence", 0.5),
            aliases=node.get("aliases", []),
            metadata=node.get("metadata", {}),
        )

    def _relation_from_edge(self, source: str, target: str, relation_data: dict):
        return Relation(
            source_id=source,
            target_id=target,
            relation_type=relation_data.get("relation_type", "related_to"),
            evidence=relation_data.get("evidence", ""),
            source_document_id=relation_data.get("source_document_id", ""),
            confidence=relation_data.get("confidence", 0.5),
            metadata=relation_data.get("metadata", {}),
        )

    def export_summary(self) -> dict:
        return {
            "nodes": self.graph.graph.number_of_nodes(),
            "edges": self.graph.graph.number_of_edges(),
            "entities": len(self.graph.entities()),
            "relations": len(self.graph.relations()),
            "neo4j_enabled": self.neo4j_enabled,
            "code_graph_enabled": self.code_graph.enabled,
        }

    def close(self) -> None:
        if self.neo4j_store:
            self.neo4j_store.close()
